# -*- coding: utf-8 -*-
"""Le a planilha de dados brutos e CALCULA os indicadores de cada slide.

Entrada: Modelo_Dados_RH.xlsx (abas Config, Base_Colaboradores, Historico_Mensal,
Custo_Rescisorio, Absenteismo, Terceiros, Inf_Adicionais).

Saida: montar_slides() devolve (meta, slides) no mesmo formato que o renderizador
do gerar_apresentacao.py espera. As posicoes/tipo de grafico/titulos vem do
spec.py; as LINHAS (rows) sao calculadas aqui a partir do periodo escolhido.
"""
import datetime as _dt
import unicodedata
from collections import OrderedDict, defaultdict

import openpyxl

from . import spec
from .periodo import Periodo, parse_mes, label_mes, MESES_PT


# ===========================================================================
# Helpers de parsing
# ===========================================================================
def _norm(s):
    if s is None:
        return ""
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    return "".join(ch for ch in s.lower() if ch.isalnum())


def _to_date(v):
    if v is None or v == "":
        return None
    if isinstance(v, _dt.datetime):
        return v.date()
    if isinstance(v, _dt.date):
        return v
    s = str(v).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y", "%m/%d/%Y"):
        try:
            return _dt.datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _to_num(v):
    if isinstance(v, (int, float)):
        return float(v)
    if v is None:
        return 0.0
    s = str(v).strip().replace("R$", "").replace("%", "").replace(" ", "")
    if not s:
        return 0.0
    # remove separador de milhar e troca decimal
    if s.count(",") and s.count("."):
        s = s.replace(".", "").replace(",", ".")
    elif s.count(","):
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return 0.0


def _bool(v):
    return _norm(v) in ("s", "sim", "true", "1", "x", "verdadeiro")


def _genero(v):
    n = _norm(v)
    if n.startswith("f"):
        return "F"
    if n.startswith("m"):
        return "M"
    return ""


def _motivo(v):
    n = _norm(v)
    if n.startswith("vol"):
        return "Voluntario"
    if n.startswith("invol"):
        return "Involuntario"
    return ""


# ===========================================================================
# Leitura das abas
# ===========================================================================
def _sheet_rows(wb, *aliases):
    """Devolve lista de dicts {col_normalizada: valor} para a 1a aba que casar."""
    target = {_norm(a) for a in aliases}
    ws = None
    for name in wb.sheetnames:
        if _norm(name) in target or any(_norm(name).startswith(t) for t in target):
            ws = wb[name]
            break
    if ws is None:
        return []
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [_norm(h) for h in rows[0]]
    out = []
    for r in rows[1:]:
        if all(c is None or str(c).strip() == "" for c in r):
            continue
        d = {headers[i]: r[i] for i in range(min(len(headers), len(r)))}
        out.append(d)
    return out


def _first(d, *aliases):
    for a in aliases:
        k = _norm(a)
        if k in d and d[k] not in (None, ""):
            return d[k]
    return None


def _load_base(wb):
    base = []
    for d in _sheet_rows(wb, "Base_Colaboradores", "Base", "Colaboradores"):
        nome = _first(d, "nome")
        if _norm(_first(d, "matricula") or nome).startswith("exemplo"):
            pass  # mantemos exemplo; gerente apaga manualmente
        c = {
            "matricula": _first(d, "matricula"),
            "nome": nome,
            "genero": _genero(_first(d, "genero", "sexo")),
            "nascimento": _to_date(_first(d, "datanascimento", "nascimento")),
            "admissao": _to_date(_first(d, "dataadmissao", "admissao")),
            "desligamento": _to_date(_first(d, "datadesligamento", "desligamento")),
            "motivo": _motivo(_first(d, "motivodesligamento", "motivo")),
            "situacao": _first(d, "situacao"),
            "vinculo": _first(d, "tipovinculo", "vinculo") or "CLT",
            "nivel": _first(d, "nivel", "classefuncional", "classe") or "Nao Inf.",
            "cargo": _first(d, "cargo"),
            "lider": _bool(_first(d, "lider")),
            "negocio": _first(d, "negocio") or "Nao Inf.",
            "diretoria": _first(d, "diretoria") or "Nao Inf.",
            "gerencia": _first(d, "gerencia") or "Nao Inf.",
            "raca": _first(d, "raca", "cor") or "Nao Informado",
            "pcd": _bool(_first(d, "pcd")),
        }
        if c["admissao"] is None and c["matricula"] is None:
            continue
        base.append(c)
    return base


def _load_historico(wb):
    hist = {}
    for d in _sheet_rows(wb, "Historico_Mensal", "Historico"):
        ano = _first(d, "ano")
        mes = parse_mes(_first(d, "mes"))
        if ano is None or mes is None:
            continue
        hist[(int(ano), int(mes))] = d
    return hist


def _load_lancamentos(wb, *aliases):
    out = []
    for d in _sheet_rows(wb, *aliases):
        ano = _first(d, "ano")
        mes = parse_mes(_first(d, "mes"))
        if ano is None or mes is None:
            continue
        d["_ano"], d["_mes"] = int(ano), int(mes)
        out.append(d)
    return out


def _load_terceiros(wb):
    """Cadastro de terceiros (1 linha por pessoa) com datas de entrada/saida."""
    out = []
    for d in _sheet_rows(wb, "Terceiros", "Terceiro"):
        nome = _first(d, "nome")
        entrada = _to_date(_first(d, "dataentrada", "entrada"))
        if nome is None and entrada is None:
            continue
        out.append({
            "nome": nome,
            "cargo": _first(d, "cargo"),
            "genero": _genero(_first(d, "genero", "sexo")),
            "entrada": entrada,
            "saida": _to_date(_first(d, "datasaida", "saida")),
            "situacao": _first(d, "situacao"),
            "parceiro": _first(d, "parceiro", "fornecedor") or "Nao Inf.",
            "diretoria": _first(d, "diretoria") or "Nao Inf.",
            "gerencia": _first(d, "gerencia") or "Nao Inf.",
            "projeto": _first(d, "projeto"),
            "orcamento": (_first(d, "orcamento", "orcamento") or "Nao Inf."),
        })
    return out


def _to_horas(v):
    """Converte horas para float; aceita timedelta (dias+horas) ou numero."""
    if isinstance(v, _dt.timedelta):
        return v.total_seconds() / 3600.0
    return _to_num(v)


# ===========================================================================
# Helpers de dominio
# ===========================================================================
def _anos(d1, d2):
    return (d2 - d1).days / 365.25


def faixa_tempo(anos):
    bins = [(0, 5, "0-05"), (5, 10, "06-10"), (10, 15, "11-15"), (15, 20, "16-20"),
            (20, 25, "21-25"), (25, 30, "26-30"), (30, 35, "31-35"),
            (35, 40, "36-40")]
    for lo, hi, lab in bins:
        if lo <= anos < hi:
            return lab
    return "41+"


FAIXAS_TEMPO = ["0-05", "06-10", "11-15", "16-20", "21-25", "26-30", "31-35",
                "36-40", "41+"]


def faixa_idade(anos):
    if anos <= 20:
        return "0-20"
    if anos <= 30:
        return "21-30"
    if anos <= 40:
        return "31-40"
    if anos <= 50:
        return "41-50"
    return "51+"


FAIXAS_IDADE = ["0-20", "21-30", "31-40", "41-50", "51+"]


def geracao(nasc):
    if nasc is None:
        return "Nao Inf."
    y = nasc.year
    if y <= 1964:
        return "Baby Boomers"
    if y <= 1980:
        return "Geracao X"
    if y <= 1996:
        return "Geracao Y"
    return "Geracao Z"


def _prefix_num(s):
    s = str(s).strip()
    num = ""
    for ch in s:
        if ch.isdigit():
            num += ch
        else:
            break
    return int(num) if num else 999


def _pct(n, d):
    return (100.0 * n / d) if d else 0.0


def _pct_str(n, d):
    return ("%.2f%%" % _pct(n, d)).replace(".", ",")


def _count_by(items, keyfn, order=None):
    c = defaultdict(int)
    for it in items:
        c[keyfn(it)] += 1
    if order:
        return [(k, c.get(k, 0)) for k in order if c.get(k, 0) or k in order]
    return sorted(c.items(), key=lambda kv: (-kv[1], str(kv[0])))


# ===========================================================================
# Contexto de calculo
# ===========================================================================
class Contexto:
    def __init__(self, wb):
        self.config = _load_config(wb)
        self.periodo = _periodo_from_config(self.config)
        self.base = _load_base(wb)
        self.historico = _load_historico(wb)
        self.resc = _load_lancamentos(wb, "Custo_Rescisorio", "CustoRescisorio")
        self.absenteismo = _load_lancamentos(wb, "Absenteismo")
        self.terceiros = _load_terceiros(wb)
        self.inf = _load_lancamentos(wb, "Inf_Adicionais", "InfAdicionais",
                                     "InformacoesAdicionais")
        self.turnover_formula = (self.config.get("formula_turnover") or "media").lower()

    # ---- subconjuntos ----
    def ativos(self, fim=None):
        fim = fim or self.periodo.fim
        return [c for c in self.base
                if c["admissao"] and c["admissao"] <= fim
                and (c["desligamento"] is None or c["desligamento"] > fim)]

    def admitidos(self, ini=None, fim=None):
        ini = ini or self.periodo.inicio
        fim = fim or self.periodo.fim
        return [c for c in self.base if c["admissao"] and ini <= c["admissao"] <= fim]

    def desligados(self, ini=None, fim=None):
        ini = ini or self.periodo.inicio
        fim = fim or self.periodo.fim
        return [c for c in self.base
                if c["desligamento"] and ini <= c["desligamento"] <= fim]

    # ---- serie mensal (historico OU derivado da base) ----
    def serie_mes(self, ano, mes):
        h = self.historico.get((ano, mes))
        ini = _dt.date(ano, mes, 1)
        fim = (_dt.date(ano, 12, 31) if mes == 12
               else _dt.date(ano, mes + 1, 1) - _dt.timedelta(days=1))
        if h is not None:
            adm = _to_num(_first(h, "admitidos"))
            desl = _to_num(_first(h, "desligados"))
            hc = _to_num(_first(h, "hc")) or len(self.ativos(fim))
            vol = _to_num(_first(h, "deslvoluntario", "voluntario"))
            invol = _to_num(_first(h, "deslinvoluntario", "involuntario"))
            span = _to_num(_first(h, "span"))
            custo = _to_num(_first(h, "custorescisorio", "custo"))
            tov = _first(h, "turnover", "turnover%")
            txd = _first(h, "txdeslig", "txdeslig%")
            turnover = _to_num(tov) if tov not in (None, "") else None
            txdeslig = _to_num(txd) if txd not in (None, "") else None
        else:
            ativos = self.ativos(fim)
            hc = len(ativos)
            adm = len(self.admitidos(ini, fim))
            d = self.desligados(ini, fim)
            desl = len(d)
            vol = sum(1 for x in d if x["motivo"] == "Voluntario")
            invol = sum(1 for x in d if x["motivo"] == "Involuntario")
            lid = sum(1 for x in ativos if x["lider"])
            span = ((hc - lid) / lid) if lid else 0
            custo = sum(_to_num(_first(r, "custo", "valor"))
                        for r in self.resc if (r["_ano"], r["_mes"]) == (ano, mes))
            turnover = txdeslig = None
        if turnover is None:
            turnover = _pct((adm + desl) / 2.0, hc)
        if txdeslig is None:
            txdeslig = _pct(desl, hc)
        return {"hc": hc, "adm": adm, "desl": desl, "vol": vol, "invol": invol,
                "span": span, "custo": custo,
                "turnover": turnover, "txdeslig": txdeslig}


def _load_config(wb):
    cfg = dict(spec.META)
    rows = _sheet_rows(wb, "Config", "Capa")
    # Config eh chave/valor: tentamos via coluna tecnica (chave) ou rotulos
    ws = None
    for name in wb.sheetnames:
        if _norm(name).startswith("config") or _norm(name).startswith("00") \
                or "capa" in _norm(name):
            ws = wb[name]
            break
    if ws is None:
        return cfg
    rotulo_map = {
        "titulo": "titulo", "tituloprincipal": "titulo",
        "subtitulo": "subtitulo", "subtituloarea": "subtitulo", "area": "area",
        "arearesponsavel": "area",
        "periodo": "periodo_label", "datadereferencia": "data_referencia",
        "datareferencia": "data_referencia",
        "anodereferencia": "ano", "ano": "ano",
        "mesdereferencia": "mes", "mes": "mes",
        "tipodeperiodo": "tipo", "tipoperiodo": "tipo",
        "janeladeevolucaomeses": "janela", "janelaevolucao": "janela", "janela": "janela",
        "formuladeturnover": "formula_turnover", "formulaturnover": "formula_turnover",
    }
    for row in ws.iter_rows(values_only=True):
        if not row:
            continue
        rot = _norm(row[0]) if len(row) > 0 else ""
        val = row[1] if len(row) > 1 else None
        # chave tecnica na coluna Z (indice 25) tem prioridade
        tech = row[25] if len(row) > 25 else None
        key = None
        if tech:
            key = str(tech)
        elif rot in rotulo_map:
            key = rotulo_map[rot]
        if key and val not in (None, ""):
            cfg[key] = val
    return cfg


def _periodo_from_config(cfg):
    ano = cfg.get("ano") or _dt.date.today().year
    mes = parse_mes(cfg.get("mes")) or _dt.date.today().month
    tipo = cfg.get("tipo") or "Mensal"
    janela = cfg.get("janela") or 12
    try:
        return Periodo(tipo, int(ano), int(mes), int(_to_num(janela) or 12))
    except Exception:
        return Periodo("Mensal", _dt.date.today().year, _dt.date.today().month)


# ===========================================================================
# Calculo por slide  (indices alinhados a spec.SLIDES)
# ===========================================================================
def _is_aprendiz(c):
    return _norm(c.get("vinculo")).startswith("aprend")


def s02_negocio(C):
    ativos = C.ativos()
    adm = C.admitidos()
    desl = C.desligados()
    negs = OrderedDict()
    for c in ativos:
        negs.setdefault(c["negocio"], 0)
    rows = []
    for neg in list(negs.keys()) + ["__total__"]:
        if neg == "__total__":
            hc = len(ativos); a = len(adm); d = len(desl)
            pcd = sum(1 for x in ativos if x["pcd"])
            apr = sum(1 for x in ativos if _is_aprendiz(x))
            rows.append(["Total", hc, a, d, pcd, apr])
        else:
            av = [x for x in ativos if x["negocio"] == neg]
            hc = len(av)
            a = sum(1 for x in adm if x["negocio"] == neg)
            d = sum(1 for x in desl if x["negocio"] == neg)
            pcd = sum(1 for x in av if x["pcd"])
            apr = sum(1 for x in av if _is_aprendiz(x))
            rows.append([neg, hc, a, d, pcd, apr])
    return {"geral": rows}


def s03_hc(C):
    ativos = C.ativos()
    fim = C.periodo.fim
    hc = len(ativos)
    masc = sum(1 for c in ativos if c["genero"] == "M")
    fem = sum(1 for c in ativos if c["genero"] == "F")
    pcd = sum(1 for c in ativos if c["pcd"])
    aprend = sum(1 for c in ativos if _norm(c["vinculo"]).startswith("aprend"))
    estag = sum(1 for c in ativos if _norm(c["vinculo"]).startswith("estag"))
    inativos = sum(1 for c in C.base
                   if c["desligamento"] and c["desligamento"] <= fim)
    kpis = [["Total Colaboradores", hc], ["Masculino", masc], ["Feminino", fem],
            ["Ativos", hc], ["Estagiarios", estag], ["PCD", pcd], ["Aprendiz", aprend]]

    piramide = _count_by(ativos, lambda c: c["nivel"])
    piramide.sort(key=lambda kv: _prefix_num(kv[0]))
    pir_rows = [[k, v] for k, v in piramide]

    tc = defaultdict(int)
    for c in ativos:
        tc[faixa_tempo(_anos(c["admissao"], fim))] += 1
    tc_rows = [[f, round(_pct(tc.get(f, 0), hc), 2)] for f in FAIXAS_TEMPO]

    ig_f = defaultdict(int); ig_m = defaultdict(int)
    for c in ativos:
        if not c["nascimento"]:
            continue
        f = faixa_idade(_anos(c["nascimento"], fim))
        if c["genero"] == "F":
            ig_f[f] += 1
        elif c["genero"] == "M":
            ig_m[f] += 1
    ig_rows = [[f, round(_pct(ig_f.get(f, 0), hc), 2),
               round(_pct(ig_m.get(f, 0), hc), 2)] for f in FAIXAS_IDADE]

    return {"kpis": kpis, "piramide": pir_rows, "tempo_casa": tc_rows,
            "idade_genero": ig_rows}


def s04_hc_evo(C):
    p = C.periodo
    series = [C.serie_mes(a, m) for (a, m) in p.meses_evo]
    labels = p.labels_evo
    # HC Realizado x Projetado (projetado vem do Historico_Mensal; se vazio, usa o realizado)
    evo = []
    for i, (a, m) in enumerate(p.meses_evo):
        real = series[i]["hc"]
        h = C.historico.get((a, m))
        proj = _to_num(_first(h, "hcprojetado", "hcprevisto", "projetado")) if h else 0
        evo.append([labels[i], real, proj or real])
    admdem = [[labels[i], series[i]["adm"], series[i]["desl"]]
              for i in range(len(labels))]
    tot_adm = sum(s["adm"] for s in series)
    tot_desl = sum(s["desl"] for s in series)
    kpis = [[f"Admitidos ({len(labels)}m)", int(tot_adm)],
            [f"Desligados ({len(labels)}m)", -int(tot_desl)],
            ["Diferenca", int(tot_adm - tot_desl)]]
    return {"kpis": kpis, "evolucao_hc": evo, "adm_dem": admdem}


def s05_admitidos(C):
    adm = C.admitidos()
    n = len(adm)
    clt = sum(1 for c in adm if _norm(c["vinculo"]).startswith(("clt", "presen")))
    ea = sum(1 for c in adm if _norm(c["vinculo"]).startswith(("estag", "aprend")))
    kpis = [["Admitidos no Periodo", n], ["CLT", clt], ["Estagio/Aprendiz", ea]]
    genero = [["Feminino", sum(1 for c in adm if c["genero"] == "F")],
              ["Masculino", sum(1 for c in adm if c["genero"] == "M")]]
    classe = _count_by(adm, lambda c: c["nivel"])
    classe.sort(key=lambda kv: _prefix_num(kv[0]))
    classe_rows = [[k, v] for k, v in classe]
    ger = _count_by(adm, lambda c: c["gerencia"])
    pareto, acc = [], 0
    for k, v in ger[:8]:
        acc += v
        pareto.append([k, v, ("%d%%" % round(_pct(acc, n)))])
    return {"kpis": kpis, "por_genero": genero, "por_classe": classe_rows,
            "pareto": pareto}


def s06_turnover(C):
    ativos = C.ativos(); hc = len(ativos)
    desl = C.desligados(); d = len(desl)
    adm = len(C.admitidos())
    vol = sum(1 for x in desl if x["motivo"] == "Voluntario")
    invol = sum(1 for x in desl if x["motivo"] == "Involuntario")
    turnover = _pct((adm + d) / 2.0, hc)
    kpis = [["Desligados", d], ["Voluntario", vol], ["Involuntario", invol],
            ["Turnover Geral", ("%.2f%%" % turnover).replace(".", ",")],
            ["Tx Deslig", _pct_str(d, hc)],
            ["Tx Vol", _pct_str(vol, hc)], ["Tx Invol", _pct_str(invol, hc)]]
    fim = C.periodo.fim
    tc = defaultdict(int)
    for c in desl:
        anos = _anos(c["admissao"], c["desligamento"]) if c["admissao"] else 0
        tc["Ate 6 Meses" if anos < 0.5 else "Acima de 2 Anos" if anos >= 2
           else "6 Meses a 2 Anos"] += 1
    tc_rows = [[k, tc.get(k, 0)]
               for k in ["Ate 6 Meses", "6 Meses a 2 Anos", "Acima de 2 Anos"]]
    ger = _count_by(desl, lambda c: geracao(c["nascimento"]),
                    order=["Baby Boomers", "Geracao X", "Geracao Y", "Geracao Z"])
    ger_rows = [[k, v] for k, v in ger if v]
    genero = [["Feminino", _pct(sum(1 for c in desl if c["genero"] == "F"), d)],
              ["Masculino", _pct(sum(1 for c in desl if c["genero"] == "M"), d)]]
    return {"kpis": kpis, "tempo_casa": tc_rows, "geracao": ger_rows,
            "genero": genero}


def _serie_turnover(C, meses):
    s = [C.serie_mes(a, m) for (a, m) in meses]
    labels = [label_mes(a, m) for (a, m) in meses]
    rows = [[labels[i], round(s[i]["turnover"], 2), round(s[i]["txdeslig"], 2)]
            for i in range(len(labels))]
    tov_ac = sum(x["turnover"] for x in s)
    txd_ac = sum(x["txdeslig"] for x in s)
    return rows, tov_ac, txd_ac


def s07_turnover_evo(C):
    rows, tov, txd = _serie_turnover(C, C.periodo.meses_evo)
    # ano corrente (jan ate o mes de referencia) consolidado no mesmo slide
    ano = C.periodo.ano
    meses_ano = [(ano, m) for m in range(1, C.periodo.mes + 1)]
    _, tov_ano, _ = _serie_turnover(C, meses_ano)
    kpis = [["Turnover Acumulado", ("%.2f%%" % tov).replace(".", ",")],
            ["Tx Deslig Acumulado", ("%.2f%%" % txd).replace(".", ",")],
            ["Turnover Ano Corrente", ("%.2f%%" % tov_ano).replace(".", ",")]]
    return {"kpis": kpis, "serie": rows}


def s09_desligados(C):
    desl = C.desligados(); d = len(desl)
    vol = sum(1 for x in desl if x["motivo"] == "Voluntario")
    invol = sum(1 for x in desl if x["motivo"] == "Involuntario")
    kpis = [["Demitidos", d], ["Voluntario", vol], ["Involuntario", invol]]
    genero = [["Feminino", sum(1 for c in desl if c["genero"] == "F")],
              ["Masculino", sum(1 for c in desl if c["genero"] == "M")],
              ["Total", d]]
    vinc = _count_by(desl, lambda c: c["vinculo"])
    vinc_rows = [[k, v] for k, v in vinc] + [["Total", d]]
    classe = _count_by(desl, lambda c: c["nivel"])
    classe.sort(key=lambda kv: _prefix_num(kv[0]))
    classe_rows = [[k, v] for k, v in classe]
    tipo = [["Voluntario", vol], ["Involuntario", invol]]
    return {"kpis": kpis, "por_genero": genero, "por_vinculo": vinc_rows,
            "por_classe": classe_rows, "por_tipo": tipo}


def s10_desligados_evo(C):
    p = C.periodo
    s = [C.serie_mes(a, m) for (a, m) in p.meses_evo]
    labels = p.labels_evo
    serie = [[labels[i], s[i]["vol"], s[i]["invol"]] for i in range(len(labels))]
    desl = C.desligados()
    by_ger = defaultdict(lambda: [0, 0])  # vol, total
    for c in desl:
        by_ger[c["gerencia"]][1] += 1
        if c["motivo"] == "Voluntario":
            by_ger[c["gerencia"]][0] += 1
    vol_rows = sorted(([g, _pct_str(v[0], v[1])] for g, v in by_ger.items()),
                      key=lambda r: r[1], reverse=True)[:6]
    ativos_ger = _count_by(C.ativos(), lambda c: c["gerencia"])
    ativos_map = dict(ativos_ger)
    tx_rows = sorted(([g, _pct_str(v[1], ativos_map.get(g, v[1]))]
                      for g, v in by_ger.items()),
                     key=lambda r: r[1], reverse=True)[:6]
    return {"serie": serie, "maiores_vol": vol_rows, "maiores_tx": tx_rows}


def s11_span(C):
    ativos = C.ativos(); hc = len(ativos)
    lid = sum(1 for c in ativos if c["lider"])
    base = hc - lid
    span = (base / lid) if lid else 0
    kpis = [["Colaboradores Ativos", hc], ["Lideranca", lid], ["Base", base],
            ["Span", ("%.2f" % span).replace(".", ",")]]
    categoria = [["Lideranca", _pct(lid, hc)], ["Base", _pct(base, hc)]]
    by_dir = defaultdict(lambda: [0, 0])  # total, lider
    for c in ativos:
        by_dir[c["diretoria"]][0] += 1
        if c["lider"]:
            by_dir[c["diretoria"]][1] += 1
    dir_rows = []
    for g, v in sorted(by_dir.items(), key=lambda kv: -kv[1][0])[:8]:
        sp = ((v[0] - v[1]) / v[1]) if v[1] else 0
        dir_rows.append([g, v[0], v[1], ("%.2f" % sp).replace(".", ",")])
    return {"kpis": kpis, "categoria": categoria, "por_diretoria": dir_rows}


PCD_COTA_PCT = 0.02   # Lei 8.213/91: 2% (empresas 100-200 func.) ate 5%
APRENDIZ_MIN = 1      # cota informada pelo RH: "de 1 a 3"
APRENDIZ_MAX = 3


def _cota_rows(status, linhas):
    """Monta as rows do bloco 'cota': 1a linha = marcador de status."""
    return [["__status__", "dentro" if status else "fora"]] + linhas


def s13_pcd(C):
    ativos = C.ativos(); hc = len(ativos)
    pcd = [c for c in ativos if c["pcd"]]
    npcd = len(pcd)
    kpis = [["Total Colaboradores", hc], ["PCD Total", npcd],
            ["% do Quadro", _pct_str(npcd, hc)]]
    exigido = int(round(PCD_COTA_PCT * hc))
    dentro = npcd >= exigido
    conformidade = _cota_rows(dentro, [
        ["Cota legal (2%)", exigido],
        ["PCD ativos", npcd],
        ["Situacao", "Dentro da cota" if dentro else "Abaixo da cota"],
    ])
    by_neg = _count_by(pcd, lambda c: c["negocio"])
    neg_rows = [[k, v, _pct_str(v, npcd)] for k, v in by_neg]
    rg_f = defaultdict(int); rg_m = defaultdict(int)
    for c in pcd:
        (rg_f if c["genero"] == "F" else rg_m)[c["raca"]] += 1
    racas = list(dict.fromkeys([c["raca"] for c in pcd]))
    rg_rows = [[r, round(_pct(rg_f.get(r, 0), npcd), 1),
               round(_pct(rg_m.get(r, 0), npcd), 1)] for r in racas]
    return {"kpis": kpis, "conformidade": conformidade, "por_negocio": neg_rows,
            "raca_genero": rg_rows}


def s_aprendiz(C):
    ativos = C.ativos(); hc = len(ativos)
    apr = [c for c in ativos if _is_aprendiz(c)]
    napr = len(apr)
    cota_txt = "de %d a %d" % (APRENDIZ_MIN, APRENDIZ_MAX)
    kpis = [["Total Colaboradores", hc], ["Aprendizes Ativos", napr],
            ["Cota Legal", cota_txt]]
    dentro = napr >= APRENDIZ_MIN
    if napr > APRENDIZ_MAX:
        situacao = "Acima da cota"
    elif dentro:
        situacao = "Dentro da cota"
    else:
        situacao = "Abaixo da cota"
    conformidade = _cota_rows(dentro, [
        ["Cota legal", cota_txt],
        ["Aprendizes ativos", napr],
        ["Situacao", situacao],
    ])
    by_dir = _count_by(apr, lambda c: c["diretoria"])
    dir_rows = [[k, v] for k, v in by_dir] or [["Sem aprendizes", 0]]
    genero = [["Feminino", sum(1 for c in apr if c["genero"] == "F")],
              ["Masculino", sum(1 for c in apr if c["genero"] == "M")]]
    return {"kpis": kpis, "conformidade": conformidade, "por_negocio": dir_rows,
            "por_genero": genero}


def s14_genero(C):
    ativos = C.ativos(); hc = len(ativos)
    masc = sum(1 for c in ativos if c["genero"] == "M")
    fem = sum(1 for c in ativos if c["genero"] == "F")
    lid = [c for c in ativos if c["lider"]]
    lid_f = sum(1 for c in lid if c["genero"] == "F")
    lid_m = sum(1 for c in lid if c["genero"] == "M")
    nlid = len(lid)
    kpis = [["Total Colaboradores", hc], ["Masculino", masc], ["Feminino", fem],
            ["% Mulheres Lideranca", _pct_str(lid_f, nlid)],
            ["% Homens Lideranca", _pct_str(lid_m, nlid)]]
    # piramide por nivel x genero
    niveis = sorted({c["nivel"] for c in ativos}, key=_prefix_num)
    pir_rows = []
    for nv in niveis:
        f = sum(1 for c in ativos if c["nivel"] == nv and c["genero"] == "F")
        m = sum(1 for c in ativos if c["nivel"] == nv and c["genero"] == "M")
        pir_rows.append([nv, f, m, f + m])
    pir_rows.append(["Total", fem, masc, hc])
    rg_f = defaultdict(int); rg_m = defaultdict(int)
    for c in ativos:
        (rg_f if c["genero"] == "F" else rg_m)[c["raca"]] += 1
    racas = list(dict.fromkeys([c["raca"] for c in ativos]))
    rg_rows = [[r, round(_pct(rg_f.get(r, 0), hc), 2),
               round(_pct(rg_m.get(r, 0), hc), 2)] for r in racas]
    lid_rows = [["Lideranca", lid_f, lid_m,
                 _pct_str(lid_f, nlid), _pct_str(lid_m, nlid)]]
    return {"kpis": kpis, "piramide": pir_rows, "raca_genero": rg_rows,
            "lideranca": lid_rows}


def s_absenteismo(C):
    """Absenteismo % = Horas Perdidas / Horas Previstas, por mes da janela."""
    by_mes = {}
    for r in C.absenteismo:
        prev = _to_horas(_first(r, "horasprevistas", "horasnormaisprogramadas",
                                "previstas"))
        perd = _to_horas(_first(r, "horasperdidas", "tempototalausencia",
                                "ausencia", "perdidas"))
        by_mes[(r["_ano"], r["_mes"])] = (prev, perd)
    serie, tot_prev, tot_perd = [], 0.0, 0.0
    for (a, m) in C.periodo.meses_evo:
        prev, perd = by_mes.get((a, m), (0.0, 0.0))
        pct = (100.0 * perd / prev) if prev else 0.0
        serie.append([label_mes(a, m), round(pct, 2)])
        tot_prev += prev; tot_perd += perd
    medio = (100.0 * tot_perd / tot_prev) if tot_prev else 0.0
    kpis = [["Absenteismo Medio", ("%.2f%%" % medio).replace(".", ",")],
            ["Horas Previstas", _milhar(tot_prev)],
            ["Horas Perdidas", _milhar(tot_perd)]]
    return {"kpis": kpis, "serie": serie}


def _no_periodo(C, lanc):
    meses = set(C.periodo.meses)
    return [r for r in lanc if (r["_ano"], r["_mes"]) in meses]


def s17_custo_rescisorio(C):
    reg_ano = [r for r in C.resc if r["_ano"] == C.periodo.ano]
    reg_per = _no_periodo(C, C.resc)
    custo_ano = sum(_to_num(_first(r, "custo", "valor")) for r in reg_ano)
    custo_per = sum(_to_num(_first(r, "custo", "valor")) for r in reg_per)
    kpis = [["Custo Resc. Ano", _mi(custo_ano)], ["Custo Resc. Periodo", _mi(custo_per)],
            ["Desligados (Periodo)", len(reg_per)]]
    # por tempo de casa (faixa via base pela matricula)
    base_map = {str(c["matricula"]): c for c in C.base if c["matricula"] is not None}
    faixa_q = defaultdict(lambda: [0, 0.0])
    for r in reg_per:
        c = base_map.get(str(_first(r, "matricula")))
        if c and c["admissao"] and c["desligamento"]:
            fx = faixa_tempo(_anos(c["admissao"], c["desligamento"]))
        else:
            fx = "Nao Inf."
        faixa_q[fx][0] += 1
        faixa_q[fx][1] += _to_num(_first(r, "custo", "valor"))
    tempo_rows = []
    tot_q = 0; tot_v = 0.0
    for fx in FAIXAS_TEMPO + ["Nao Inf."]:
        if fx in faixa_q:
            q, v = faixa_q[fx]
            tot_q += q; tot_v += v
            tempo_rows.append([fx, q, _milhar(v), _milhar(v / q if q else 0)])
    tempo_rows.append(["Total", tot_q, _milhar(tot_v),
                       _milhar(tot_v / tot_q if tot_q else 0)])
    evo = [[label_mes(a, m), round(C.serie_mes(a, m)["custo"] / 1e6, 2)]
           for (a, m) in C.periodo.meses_evo]
    return {"kpis": kpis, "por_tempo": tempo_rows, "evolucao": evo}


def s_terceiros(C):
    """Grau de terceirizacao = terceiros / (proprios ativos + terceiros)."""
    fim = C.periodo.fim
    ativos_ter = [t for t in C.terceiros
                  if t["entrada"] and t["entrada"] <= fim
                  and (t["saida"] is None or t["saida"] > fim)]
    entradas = sum(1 for t in C.terceiros
                   if t["entrada"] and C.periodo.contem(t["entrada"]))
    saidas = sum(1 for t in C.terceiros
                 if t["saida"] and C.periodo.contem(t["saida"]))
    nter = len(ativos_ter)
    proprios = len(C.ativos(fim))
    gter = _pct(nter, proprios + nter)
    kpis = [["Terceiros Ativos", nter], ["Entradas", entradas], ["Saidas", saidas],
            ["Grau Terceirizacao", ("%.1f%%" % gter).replace(".", ",")]]
    parceiro = _count_by(ativos_ter, lambda t: t["parceiro"])
    parc_rows = [[k, v] for k, v in parceiro][:8] or [["Sem terceiros", 0]]
    # GTER por diretoria: proprios x terceiros
    prop_dir = defaultdict(int)
    for c in C.ativos(fim):
        prop_dir[c["diretoria"]] += 1
    ter_dir = defaultdict(int)
    for t in ativos_ter:
        ter_dir[t["diretoria"]] += 1
    dir_rows = []
    for d in sorted(set(list(prop_dir) + list(ter_dir)),
                    key=lambda d: -ter_dir.get(d, 0)):
        p = prop_dir.get(d, 0); te = ter_dir.get(d, 0)
        dir_rows.append([d, p, te, ("%.1f%%" % _pct(te, p + te)).replace(".", ",")])
    dir_rows = dir_rows[:8] or [["Sem dados", 0, 0, "0,0%"]]
    orc = _count_by(ativos_ter, lambda t: (t["orcamento"] or "Nao Inf.").strip())
    orc_rows = [[k, v] for k, v in orc] or [["Nao Inf.", 0]]
    return {"kpis": kpis, "por_parceiro": parc_rows, "por_diretoria": dir_rows,
            "por_orcamento": orc_rows}


def _milhar(v):
    return ("{:,.0f}".format(v)).replace(",", ".")


def _mi(v):
    if v >= 1e6:
        return ("%.2f Mi" % (v / 1e6)).replace(".", ",")
    return ("%.1f Mil" % (v / 1e3)).replace(".", ",")


# Ordem ALINHADA a spec.SLIDES (1 computer por slide, mesmo indice).
COMPUTERS = [s02_negocio, s03_hc, s04_hc_evo, s05_admitidos, s06_turnover,
             s07_turnover_evo, s09_desligados, s10_desligados_evo, s11_span,
             s13_pcd, s_aprendiz, s14_genero, s_absenteismo,
             s17_custo_rescisorio, s_terceiros]


# ===========================================================================
# Montagem final (dados + layout do spec)
# ===========================================================================
def montar_slides(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    C = Contexto(wb)
    meta = dict(spec.META)
    for k in ("titulo", "subtitulo", "area", "data_referencia"):
        if C.config.get(k):
            meta[k] = C.config[k]
    meta["periodo"] = C.periodo.label

    slides = []
    for i, sl in enumerate(spec.SLIDES):
        computed = {}
        if i < len(COMPUTERS):
            try:
                computed = COMPUTERS[i](C) or {}
            except Exception as e:  # degrada sem quebrar
                print(f"[aviso] slide {i+2} ({sl['title']}): {e}")
                computed = {}
        blocks = []
        for blk in sl["blocks"]:
            rows = computed.get(blk["key"])
            if rows is None:
                rows = blk["rows"]  # fallback exemplo
            blocks.append({"key": blk["key"], "title": blk["title"],
                           "render": blk["render"], "columns": blk["columns"],
                           "rows": rows, "pos": blk["pos"]})
        slides.append({"title": sl["title"], "blocks": blocks})
    return meta, slides
