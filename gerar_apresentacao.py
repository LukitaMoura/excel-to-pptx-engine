# -*- coding: utf-8 -*-
"""Le a planilha de dados brutos, CALCULA os indicadores e gera a apresentacao
de RH com identidade visual configuravel (hr_deck/brand.py).

Uso:
    py -3 gerar_apresentacao.py [dados.xlsx] [saida.pptx]

Padrao de entrada/saida:
    Modelo_Dados_RH.xlsx  ->  Apresentacao_RH.pptx

Os calculos ficam em hr_deck/dados.py; o layout (posicoes/graficos) em
hr_deck/spec.py. Se a planilha nao existir, gera o PPT a partir dos dados de
exemplo do spec (util para visualizar o layout).
"""
import sys
import os

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION

from hr_deck import brand, spec, dados

CHART_TYPES = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
    "column_stacked": XL_CHART_TYPE.COLUMN_STACKED,
    "bar_stacked": XL_CHART_TYPE.BAR_STACKED,
}
PIE_LIKE = {"pie", "doughnut"}


# ===========================================================================
# Leitura da planilha
# ===========================================================================
def _num(v):
    if isinstance(v, (int, float)):
        return v
    if v is None:
        return 0
    s = str(v).strip().replace("%", "").replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return 0


def _get(row, idx):
    return row[idx] if idx < len(row) else None


# ===========================================================================
# Helpers de desenho
# ===========================================================================
def _set_font(run, size=11, bold=False, color=brand.TEXT, name=brand.FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = name
    run.font.color.rgb = brand.rgb(color)


def add_text(slide, l, t, w, h, text, size=11, bold=False, color=brand.TEXT,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, name=brand.FONT):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Pt(2)
    tf.margin_top = tf.margin_bottom = Pt(1)
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    _set_font(r, size, bold, color, name)
    return tb


def add_rect(slide, l, t, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE,
             line_w=1.0):
    sp = slide.shapes.add_shape(shape, Inches(l), Inches(t), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = brand.rgb(fill)
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = brand.rgb(line)
        sp.line.width = Pt(line_w)
    sp.shadow.inherit = False
    return sp


def _set_shape_alpha(shape, opacity):
    """Define opacidade (0..1) do preenchimento solido via XML (a:alpha)."""
    from pptx.oxml.ns import qn
    srgb = shape.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    for tag in ("a:alpha",):
        existing = srgb.find(qn(tag))
        if existing is not None:
            srgb.remove(existing)
    alpha = srgb.makeelement(qn("a:alpha"), {"val": str(int(opacity * 100000))})
    srgb.append(alpha)


def add_image(slide, path, l, t, w=None, h=None):
    """Insere imagem se o arquivo existir; preserva proporcao se so um lado for dado."""
    if not path or not os.path.exists(path):
        return None
    kw = {}
    if w is not None:
        kw["width"] = Inches(w)
    if h is not None:
        kw["height"] = Inches(h)
    return slide.shapes.add_picture(path, Inches(l), Inches(t), **kw)


def fmt(v):
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, float):
        return ("%.2f" % v).replace(".", ",")
    return "" if v is None else str(v)


# ===========================================================================
# Renderizadores de bloco
# ===========================================================================
def render_block_title(slide, l, t, w, title):
    if not title:
        return t
    add_rect(slide, l, t + 0.02, 0.10, 0.26, fill=brand.ACCENT)  # tracinho vermelho
    add_text(slide, l + 0.18, t, w - 0.18, 0.30, title, size=12, bold=True,
             color=brand.PRIMARY, anchor=MSO_ANCHOR.MIDDLE)
    return t + 0.38


def render_kpi(slide, block, pos):
    l, t, w, h = pos
    rows = block["rows"]
    n = max(1, len(rows))
    gap = 0.16
    cw = (w - gap * (n - 1)) / n
    for i, row in enumerate(rows):
        x = l + i * (cw + gap)
        card = add_rect(slide, x, t, cw, h, fill=brand.WHITE, line=brand.BORDER,
                        shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        try:
            card.adjustments[0] = 0.06
        except Exception:
            pass
        add_rect(slide, x, t, 0.07, h, fill=brand.ACCENT)  # accent lateral
        label = fmt(_get(row, 0))
        value = fmt(_get(row, 1))
        add_text(slide, x + 0.14, t + 0.10, cw - 0.2, h * 0.45, value,
                 size=20, bold=True, color=brand.PRIMARY, anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, x + 0.14, t + h * 0.55, cw - 0.2, h * 0.40, label,
                 size=9, bold=False, color=brand.GRAY_DARK, anchor=MSO_ANCHOR.TOP)


def render_table(slide, block, pos):
    l, t, w, h = pos
    t = render_block_title(slide, l, t, w, block["title"])
    h = max(0.6, (pos[1] + pos[3]) - t)
    cols = block["columns"]
    data = block["rows"]
    nrows = len(data) + 1
    ncols = max(1, len(cols))
    gtbl = slide.shapes.add_table(nrows, ncols, Inches(l), Inches(t),
                                  Inches(w), Inches(min(h, 0.32 * nrows))).table
    # cabecalho
    for j, col in enumerate(cols):
        cell = gtbl.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = brand.rgb(brand.PRIMARY)
        cell.margin_top = cell.margin_bottom = Pt(1)
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
        r = p.add_run(); r.text = str(col)
        _set_font(r, 9, True, brand.WHITE)
    # dados
    for i, row in enumerate(data, start=1):
        for j in range(ncols):
            cell = gtbl.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = brand.rgb(brand.PRIMARY_10 if i % 2 else brand.WHITE)
            cell.margin_top = cell.margin_bottom = Pt(1)
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            r = p.add_run(); r.text = fmt(_get(row, j))
            _set_font(r, 9, False, brand.TEXT)


def render_chart(slide, block, pos):
    l, t, w, h = pos
    t = render_block_title(slide, l, t, w, block["title"])
    h = max(0.8, (pos[1] + pos[3]) - t)
    render = block["render"]
    cols = block["columns"]
    data = block["rows"]
    if not data:
        return

    cats = [fmt(_get(r, 0)) for r in data]
    cd = CategoryChartData()
    cd.categories = cats
    nseries = max(1, len(cols) - 1)
    for s in range(nseries):
        name = cols[s + 1] if s + 1 < len(cols) else f"Serie {s+1}"
        cd.add_series(name, tuple(_num(_get(r, s + 1)) for r in data))

    gf = slide.shapes.add_chart(CHART_TYPES[render], Inches(l), Inches(t),
                                Inches(w), Inches(h), cd)
    chart = gf.chart
    chart.font.name = brand.FONT
    chart.font.size = Pt(9)

    # legenda
    if render in PIE_LIKE or nseries > 1:
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(8)
    else:
        chart.has_legend = False

    plot = chart.plots[0]

    if render in PIE_LIKE:
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.show_value = False
        dl.show_percentage = True
        dl.number_format = '0.0%'
        dl.number_format_is_linked = False
        dl.font.size = Pt(8)
        if render == "pie":
            try:
                dl.position = XL_LABEL_POSITION.OUTSIDE_END
            except Exception:
                pass
        series = plot.series[0]
        for i, point in enumerate(series.points):
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = brand.rgb(
                brand.SERIES_COLORS[i % len(brand.SERIES_COLORS)])
    else:
        for idx, series in enumerate(chart.series):
            color = brand.rgb(brand.SERIES_COLORS[idx % len(brand.SERIES_COLORS)])
            if render == "line":
                series.format.line.color.rgb = color
                series.format.line.width = Pt(2.25)
            else:
                series.format.fill.solid()
                series.format.fill.fore_color.rgb = color
        try:
            chart.value_axis.has_major_gridlines = True
            chart.value_axis.major_gridlines.format.line.color.rgb = brand.rgb(brand.PRIMARY_10)
            chart.category_axis.tick_labels.font.size = Pt(8)
            chart.value_axis.tick_labels.font.size = Pt(8)
        except Exception:
            pass


def render_cota(slide, block, pos):
    """Selo de conformidade da cota legal (verde 'dentro' / vermelho 'fora').

    A 1a linha do bloco e o marcador ["__status__", "dentro"|"fora"]; as demais
    sao pares (rotulo, valor) exibidos como legenda ao lado do selo.
    """
    l, t, w, h = pos
    t = render_block_title(slide, l, t, w, block["title"])
    h = max(0.8, (pos[1] + pos[3]) - t)
    rows = list(block["rows"])
    status = "dentro"
    if rows and str(_get(rows[0], 0)) == "__status__":
        status = str(_get(rows[0], 1) or "dentro").strip().lower()
        rows = rows[1:]
    ok = status.startswith("dentro")
    cor = brand.GREEN if ok else brand.ACCENT
    cor_bg = brand.GREEN_10 if ok else brand.ACCENT_10
    simbolo = "OK" if ok else "X"
    # rotulo proeminente = texto da linha "Situacao" (se houver); senao, default
    rotulo = "Dentro da cota" if ok else "Fora da cota"
    for row in rows:
        rot = str(_get(row, 0) or "").strip().lower()
        if rot.startswith("situac") and _get(row, 1):
            rotulo = str(_get(row, 1))
            break

    # cartao com selo a esquerda + legenda a direita
    add_rect(slide, l, t, w, h, fill=cor_bg, line=brand.BORDER,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    selo = min(1.7, h - 0.6)
    add_rect(slide, l + 0.30, t + (h - selo) / 2, selo, selo, fill=cor,
             shape=MSO_SHAPE.OVAL)
    add_text(slide, l + 0.30, t + (h - selo) / 2, selo, selo, simbolo,
             size=26, bold=True, color=brand.WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    tx = l + 0.30 + selo + 0.30
    tw = w - (tx - l) - 0.25
    add_text(slide, tx, t + 0.22, tw, 0.4, rotulo, size=15, bold=True, color=cor)
    yy = t + 0.78
    for row in rows:
        rot = fmt(_get(row, 0)); val = fmt(_get(row, 1))
        add_text(slide, tx, yy, tw, 0.34,
                 f"{rot}:  {val}" if rot else val,
                 size=11, color=brand.TEXT)
        yy += 0.40


def render_block(slide, block, pos):
    r = block["render"]
    if r == "kpi":
        render_kpi(slide, block, pos)
    elif r == "cota":
        render_cota(slide, block, pos)
    elif r == "table":
        render_table(slide, block, pos)
    elif r in CHART_TYPES:
        render_chart(slide, block, pos)
    else:
        render_table(slide, block, pos)


# ===========================================================================
# Estrutura dos slides
# ===========================================================================
def add_header(slide, idx, total, title):
    add_rect(slide, 0, 0, 13.333, 0.82, fill=brand.PRIMARY)          # faixa azul
    add_rect(slide, 0, 0.82, 13.333, 0.06, fill=brand.ACCENT)        # linha vermelha
    # logo (fallback: nome da empresa em texto)
    if add_image(slide, brand.LOGO_ICON, 0.30, 0.13, h=0.56):
        title_x = 1.15
    else:
        add_text(slide, 0.4, 0, 1.6, 0.82, brand.COMPANY_SHORT, size=18, bold=True,
                 color=brand.WHITE, anchor=MSO_ANCHOR.MIDDLE)
        title_x = 2.0
    add_text(slide, title_x, 0, 10.0, 0.82, title, size=17, bold=True,
             color=brand.WHITE, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, 11.4, 0, 1.5, 0.82, f"{idx:02d} / {total:02d}", size=11,
             color=brand.WHITE, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)


def build_cover(prs, meta):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    # fundo institucional (fallback: cor primaria)
    if not add_image(slide, brand.COVER_BG, 0, 0, w=13.333, h=7.5):
        add_rect(slide, 0, 0, 13.333, 7.5, fill=brand.PRIMARY)
        add_rect(slide, 0, 5.9, 13.333, 1.6, fill=brand.ACCENT)
    # painel translucido escuro p/ garantir contraste do texto (lado esquerdo)
    panel = add_rect(slide, 0, 1.9, 8.2, 3.6, fill=brand.PRIMARY)
    _set_shape_alpha(panel, 0.62)  # 62% opaco -> deixa o fundo aparecer atras
    # wordmark (fallback: nome da empresa em texto)
    if not add_image(slide, brand.LOGO_WORD, 0.85, 0.6, h=0.9):
        add_text(slide, 0.9, 0.6, 5.0, 0.9, brand.COMPANY_SHORT, size=30, bold=True,
                 color=brand.WHITE, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, 0.9, 2.15, 7.2, 1.6, meta.get("titulo", ""), size=36, bold=True,
             color=brand.WHITE)
    add_text(slide, 0.9, 3.75, 7.2, 0.6, meta.get("subtitulo", ""), size=20,
             color=brand.WHITE, name=brand.FONT)
    add_text(slide, 0.9, 4.45, 7.2, 0.7, meta.get("periodo", ""), size=20, bold=True,
             color="FFFFFF")
    add_text(slide, 0.9, 6.5, 9.0, 0.9,
             f"{meta.get('area','')}   |   Data de referencia: {meta.get('data_referencia','')}",
             size=12, color=brand.WHITE, anchor=MSO_ANCHOR.MIDDLE)


def build_closing(prs, meta):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, 13.333, 7.5, fill=brand.PRIMARY)
    add_rect(slide, 0, 3.55, 13.333, 0.06, fill=brand.ACCENT)
    add_text(slide, 0, 2.4, 13.333, 1.0, "Obrigado", size=44, bold=True,
             color=brand.WHITE, align=PP_ALIGN.CENTER)
    add_text(slide, 0, 3.7, 13.333, 0.6, meta.get("area", ""), size=16,
             color=brand.ACCENT_40, align=PP_ALIGN.CENTER)
    # wordmark centralizado no rodape (fallback: nada)
    add_image(slide, brand.LOGO_WORD, 5.97, 5.6, h=0.7)


def build_presentation(meta, slides, out):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    build_cover(prs, meta)
    total = len(slides)
    for idx, sl in enumerate(slides, start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_header(slide, idx, total, sl["title"])
        for block in sl["blocks"]:
            render_block(slide, block, block["pos"])
    build_closing(prs, meta)

    prs.save(out)
    return prs


def slides_from_spec():
    """Fallback: usa os dados de exemplo do spec (sem planilha de dados)."""
    out = []
    for sl in spec.SLIDES:
        blocks = [dict(b) for b in sl["blocks"]]
        out.append({"title": sl["title"], "blocks": blocks})
    return out


def main():
    src = sys.argv[1] if len(sys.argv) > 1 else "Modelo_Dados_RH.xlsx"
    out = sys.argv[2] if len(sys.argv) > 2 else "Apresentacao_RH.pptx"

    if os.path.exists(src):
        meta, slides = dados.montar_slides(src)
        print(f"Lendo e calculando dados de: {src}  |  periodo: {meta.get('periodo')}")
    else:
        meta, slides = dict(spec.META), slides_from_spec()
        print(f"[aviso] '{src}' nao encontrado - usando dados de exemplo do spec.")

    build_presentation(meta, slides, out)
    print(f"OK -> {out} ({len(slides) + 2} slides, incluindo capa e fechamento)")


if __name__ == "__main__":
    main()
