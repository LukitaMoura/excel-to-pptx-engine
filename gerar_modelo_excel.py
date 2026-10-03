# -*- coding: utf-8 -*-
"""Gera a planilha-modelo de DADOS BRUTOS que a gestora preenche.

A automacao (gerar_apresentacao.py) le esta planilha, CALCULA todos os
indicadores e monta o PowerPoint, filtrando pelo periodo escolhido na aba Config.

Abas:
  Config             -> capa + painel de controle (periodo: Mensal/Semestral/Anual)
  Base_Colaboradores -> 1 linha por colaborador (nucleo dos calculos)
  Historico_Mensal   -> 1 linha por mes (opcional; sobrepoe a evolucao derivada);
                        inclui HC_Projetado p/ o grafico Realizado x Projetado
  Custo_Rescisorio   -> lancamentos (Diretoria / Area / Custo)
  Absenteismo        -> 1 linha por mes (Horas Previstas / Horas Perdidas)
  Terceiros          -> 1 linha por terceiro (entrada/saida/parceiro/orcamento)
  Inf_Adicionais     -> cota legal PCD/Aprendiz por mes (opcional)
  Leia-me            -> instrucoes + dicionario de valores

A planilha sai com um EXEMPLO sintetico (apague e cole seus dados reais).

Uso:
    py -3 gerar_modelo_excel.py
Gera: Modelo_Dados_RH.xlsx
"""
import datetime as dt
import random

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

from hr_deck import brand

OUT = "Modelo_Dados_RH.xlsx"
TECH_COL = 26  # coluna Z (oculta) para chaves tecnicas na aba Config

# estilos -------------------------------------------------------------------
F_H1   = Font(name=brand.FONT_FALLBACK, size=16, bold=True, color=brand.PRIMARY)
F_HEAD = Font(name=brand.FONT_FALLBACK, size=10, bold=True, color="FFFFFF")
F_CELL = Font(name=brand.FONT_FALLBACK, size=10, color=brand.TEXT)
F_HINT = Font(name=brand.FONT_FALLBACK, size=9, italic=True, color=brand.GRAY_DARK)

FILL_HEAD = PatternFill("solid", fgColor=brand.PRIMARY)
FILL_KEY  = PatternFill("solid", fgColor=brand.PRIMARY_80)
FILL_ALT  = PatternFill("solid", fgColor=brand.PRIMARY_10)

thin = Side(style="thin", color=brand.BORDER)
BORD = Border(left=thin, right=thin, top=thin, bottom=thin)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="center")


# ---------------------------------------------------------------------------
def _table(ws, headers, rows, widths=None, date_cols=()):
    for j, h in enumerate(headers, start=1):
        c = ws.cell(1, j, h)
        c.font, c.fill, c.alignment, c.border = F_HEAD, FILL_HEAD, CENTER, BORD
    for i, row in enumerate(rows, start=2):
        for j, val in enumerate(row, start=1):
            c = ws.cell(i, j, val)
            c.font, c.border = F_CELL, BORD
            c.alignment = LEFT if j in (1, 2) else CENTER
            if (j - 1) in date_cols and isinstance(val, dt.date):
                c.number_format = "DD/MM/YYYY"
            if i % 2 == 1:
                c.fill = FILL_ALT
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    if widths:
        for j, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(j)].width = w


def _dv(ws, col_letter, options, last=600):
    dv = DataValidation(type="list", formula1='"%s"' % ",".join(options),
                        allow_blank=True, showErrorMessage=False)
    ws.add_data_validation(dv)
    dv.add("%s2:%s%d" % (col_letter, col_letter, last))


# ---------------------------------------------------------------------------
# Exemplo sintetico
# ---------------------------------------------------------------------------
def _exemplo_base(n=220):
    random.seed(42)
    niveis = (["02 - A"] * 6 + ["03 - P"] * 7 + ["04 - E"] * 6 + ["06 - C"] * 4 +
              ["08 - G"] * 2 + ["09 - GS"] + ["10 - DIR"])
    diretorias = ["DIR Financeira", "DIR Tecnologia", "DIR Compras",
                  "DIR Recursos Humanos"]
    gerencias = ["GER Folha", "GER Sistemas", "GER Compras", "GER RH",
                 "GER Faturamento", "GER Infraestrutura"]
    racas = (["Branca"] * 5 + ["Parda"] * 3 + ["Preta"] * 1 + ["Amarela"] * 1 +
             ["Nao Informado"] * 1)
    cargo_por_nivel = {"02 - A": "Assistente", "03 - P": "Analista",
                       "04 - E": "Analista Senior", "06 - C": "Coordenador",
                       "08 - G": "Gerente", "09 - GS": "Gerente Senior",
                       "10 - DIR": "Diretor"}
    lider_niveis = {"06 - C", "08 - G", "09 - GS", "10 - DIR"}

    rows = []
    # admitidos e desligados dentro da janela (proporcionais ao tamanho da base)
    n_adm = max(3, round(n * 0.03))
    n_desl = max(12, round(n * 0.12))
    admit_fev = set(range(0, n_adm))                       # admitidos em fev/2022
    desliga_idx = set(range(n_adm, n_adm + n_desl))        # desligados na janela
    deslig_fev = set(range(n_adm, n_adm + max(3, n_desl // 3)))  # parte em fev/22

    for i in range(n):
        nivel = random.choice(niveis)
        genero = "M" if random.random() < 0.56 else "F"
        nasc = dt.date(random.randint(1962, 2001), random.randint(1, 12),
                       random.randint(1, 28))
        if i in admit_fev:
            adm = dt.date(2022, 2, random.randint(1, 18))
        else:
            adm = dt.date(random.randint(2005, 2021), random.randint(1, 12),
                          random.randint(1, 28))
        desl, motivo, situacao = None, None, "Ativo"
        if i in desliga_idx:
            if i in deslig_fev:
                desl = dt.date(2022, 2, random.randint(1, 27))
            else:
                ano = random.choice([2021, 2021, 2022])
                mes = random.randint(3, 12) if ano == 2021 else 1
                desl = dt.date(ano, mes, random.randint(1, 27))
            motivo = "Voluntario" if random.random() < 0.66 else "Involuntario"
            situacao = "Inativo"
        lider = "S" if (nivel in lider_niveis and random.random() < 0.8) else "N"
        vinculo = "CLT"
        if nivel == "02 - A" and random.random() < 0.4:
            vinculo = random.choice(["Estagio", "Aprendiz"])
        rows.append([
            10001 + i, f"Colaborador {i+1}", genero, nasc, adm, desl,
            motivo, situacao, vinculo, nivel, cargo_por_nivel[nivel], lider,
            "Suporte ao Negocio", random.choice(diretorias),
            random.choice(gerencias), random.choice(racas),
            "S" if random.random() < 0.07 else "N",
        ])
    return rows


def _exemplo_historico(hc_atual, span):
    """6 meses ficticios (set/21..fev/22), coerentes com a base.

    Termina no HC atual da base e usa volumes proporcionais ao tamanho.
    """
    meses = [(2021, 9), (2021, 10), (2021, 11), (2021, 12), (2022, 1), (2022, 2)]
    adm_pat = [0.012, 0.008, 0.018, 0.010, 0.006, 0.020]
    desl_pat = [0.015, 0.010, 0.022, 0.016, 0.011, 0.025]
    span_off = [0.1, 0.1, 0.0, 0.0, 0.0, 0.0]
    rows = []
    for k, (ano, mes) in enumerate(meses):
        hc = hc_atual + (5 - k)                       # drift suave ate o atual
        hc_proj = hc + round(hc * (0.010 - 0.002 * k))  # meta ligeiramente acima
        adm = max(1, round(hc * adm_pat[k]))
        desl = max(1, round(hc * desl_pat[k]))
        vol = round(desl * 0.66)
        invol = desl - vol
        turnover = round((adm + desl) / 2.0 / hc * 100, 2)
        txdeslig = round(desl / hc * 100, 2)
        custo = desl * 48000
        rows.append([ano, mes, hc, hc_proj, adm, desl, turnover, txdeslig, vol,
                     invol, round(span + span_off[k], 1), custo])
    return rows


def _exemplo_rescisorio(base_rows):
    """Custo rescisorio: 1 linha por desligado (Diretoria / Area / Custo)."""
    random.seed(7)
    resc = []
    for r in base_rows:
        if r[5] is not None:  # desligado -> custo rescisorio
            resc.append([r[5].year, r[5].month, r[0], r[13], r[14],
                         round(random.uniform(5000, 160000), 2)])
    return resc


def _exemplo_absenteismo(hc_atual):
    """1 linha por mes (set/21..fev/22): horas previstas x perdidas."""
    random.seed(11)
    meses = [(2021, 9), (2021, 10), (2021, 11), (2021, 12), (2022, 1), (2022, 2)]
    rows = []
    for (ano, mes) in meses:
        previstas = hc_atual * 176                      # ~176h/mes por colaborador
        taxa = random.uniform(0.006, 0.020)             # 0,6% a 2,0%
        perdidas = round(previstas * taxa, 1)
        rows.append([ano, mes, previstas, perdidas])
    return rows


def _exemplo_terceiros(n=46):
    random.seed(23)
    cargos = ["Assistente", "Analista", "Analista Senior", "Coordenador", "Gerente"]
    parceiros = ["Parceiro A", "Parceiro B", "Parceiro C", "Parceiro D"]
    diretorias = ["DIR Financeira", "DIR Tecnologia", "DIR Compras",
                  "DIR Recursos Humanos", "DIR Operacoes"]
    gerencias = ["GER Folha", "GER Sistemas", "GER Compras", "GER RH",
                 "GER Faturamento", "GER Infraestrutura"]
    projetos = ["NOC", "BPO", "Fabrica I", "Fabrica II", "Pool", "Recepcao"]
    rows = []
    for i in range(n):
        entrada = dt.date(random.randint(2011, 2022), random.randint(1, 12),
                          random.randint(1, 28))
        if random.random() < 0.55:
            saida = dt.date(random.choice([2021, 2022]), random.randint(1, 12),
                            random.randint(1, 28))
            if saida <= entrada:
                saida = None
            situacao = "Inativo" if saida else "Ativo"
        else:
            saida, situacao = None, "Ativo"
        rows.append([
            f"Terceiro {i+1}", random.choice(cargos),
            "M" if random.random() < 0.5 else "F", entrada, saida, situacao,
            parceiros[i % len(parceiros)], random.choice(diretorias),
            random.choice(gerencias), random.choice(projetos),
            random.choice(["CAPEX", "OPEX"]),
        ])
    return rows


def _exemplo_inf(base_rows):
    """Cota legal PCD/Aprendiz por mes (referencia opcional)."""
    ativos = [r for r in base_rows if r[5] is None]
    total = len(ativos)
    pcd = sum(1 for r in ativos if r[16] == "S")
    apr = sum(1 for r in ativos if str(r[8]).lower().startswith("aprend"))
    rows = []
    for (ano, mes) in [(2022, 1), (2022, 2)]:
        rows.append([ano, mes, total, pcd, round(pcd / total * 100, 2) if total else 0,
                     apr])
    return rows


# ---------------------------------------------------------------------------
# Abas
# ---------------------------------------------------------------------------
def build_config(wb):
    ws = wb.create_sheet("Config")
    ws.merge_cells("A1:C1")
    ws.cell(1, 1, "Config / Painel de controle").font = F_H1
    ws.merge_cells("A2:C2")
    ws.cell(2, 1, "Edite a coluna B. Periodo e o que define o recorte da "
                  "apresentacao (Mensal / Semestral / Anual).").font = F_HINT

    linhas = [
        ("Titulo principal", "Indicadores de Recursos Humanos", "titulo"),
        ("Subtitulo / area", "Suporte ao Negocio", "subtitulo"),
        ("Area responsavel", "Recursos Humanos - Empresa Exemplo", "area"),
        ("Data de referencia", "28/02/2022", "data_referencia"),
        ("Ano de referencia", 2022, "ano"),
        ("Mes de referencia", "fevereiro", "mes"),
        ("Tipo de periodo", "Mensal", "tipo"),
        ("Janela de evolucao (meses)", 6, "janela"),
        ("Formula de turnover", "media", "formula_turnover"),
    ]
    r = 4
    for rotulo, valor, key in linhas:
        lc = ws.cell(r, 1, rotulo)
        lc.font, lc.fill, lc.alignment, lc.border = F_HEAD, FILL_KEY, LEFT, BORD
        vc = ws.cell(r, 2, valor)
        vc.font, vc.alignment, vc.border = F_CELL, LEFT, BORD
        ws.cell(r, TECH_COL, key)  # chave tecnica oculta
        r += 1

    meses = ["janeiro", "fevereiro", "marco", "abril", "maio", "junho", "julho",
             "agosto", "setembro", "outubro", "novembro", "dezembro"]
    # validacao especifica das celulas de mes (B9) e tipo (B10)
    dv_mes = DataValidation(type="list", formula1='"%s"' % ",".join(meses),
                            allow_blank=True)
    ws.add_data_validation(dv_mes); dv_mes.add("B9")
    dv_tipo = DataValidation(type="list", formula1='"Mensal,Semestral,Anual"',
                             allow_blank=True)
    ws.add_data_validation(dv_tipo); dv_tipo.add("B10")

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 42
    ws.column_dimensions[get_column_letter(TECH_COL)].hidden = True
    ws.sheet_view.showGridLines = False


def build_base(wb, base_rows):
    ws = wb.create_sheet("Base_Colaboradores")
    headers = ["Matricula", "Nome", "Genero", "DataNascimento", "DataAdmissao",
               "DataDesligamento", "MotivoDesligamento", "Situacao", "TipoVinculo",
               "Nivel", "Cargo", "Lider", "Negocio", "Diretoria", "Gerencia",
               "Raca", "PCD"]
    widths = [11, 26, 9, 15, 14, 16, 18, 11, 12, 11, 16, 8, 20, 20, 20, 16, 7]
    _table(ws, headers, base_rows, widths, date_cols=(3, 4, 5))
    _dv(ws, "C", ["M", "F"])
    _dv(ws, "G", ["Voluntario", "Involuntario"])
    _dv(ws, "H", ["Ativo", "Afastado", "Inativo"])
    _dv(ws, "I", ["CLT", "Estagio", "Aprendiz"])
    _dv(ws, "L", ["S", "N"])
    _dv(ws, "Q", ["S", "N"])
    _dv(ws, "P", ["Branca", "Parda", "Preta", "Amarela", "Indigena",
                  "Nao Informado"])


def build_historico(wb, hist_rows):
    ws = wb.create_sheet("Historico_Mensal")
    headers = ["Ano", "Mes", "HC", "HC_Projetado", "Admitidos", "Desligados",
               "Turnover%", "TxDeslig%", "DeslVoluntario", "DeslInvoluntario",
               "Span", "CustoRescisorio"]
    # Preenchido com exemplo de 6 meses. Esta aba SOBREPOE a evolucao calculada
    # da base; deixe vazia se preferir que a evolucao seja derivada da base.
    # HC_Projetado alimenta a linha "Projetado" do grafico de evolucao de HC.
    _table(ws, headers, hist_rows, [6, 6] + [13] * 10)


def build_rescisorio(wb, resc):
    ws = wb.create_sheet("Custo_Rescisorio")
    _table(ws, ["Ano", "Mes", "Matricula", "Diretoria", "Area", "Custo"], resc,
           [6, 6, 11, 22, 20, 14])


def build_absenteismo(wb, abs_rows):
    ws = wb.create_sheet("Absenteismo")
    _table(ws, ["Ano", "Mes", "HorasPrevistas", "HorasPerdidas"], abs_rows,
           [6, 6, 16, 16])


def build_terceiros(wb, ter_rows):
    ws = wb.create_sheet("Terceiros")
    headers = ["Nome", "Cargo", "Genero", "DataEntrada", "DataSaida", "Situacao",
               "Parceiro", "Diretoria", "Gerencia", "Projeto", "Orcamento"]
    _table(ws, headers, ter_rows,
           [22, 16, 9, 14, 14, 11, 14, 22, 20, 14, 12], date_cols=(3, 4))
    _dv(ws, "C", ["M", "F"])
    _dv(ws, "F", ["Ativo", "Inativo"])
    _dv(ws, "K", ["CAPEX", "OPEX"])


def build_inf_adicionais(wb, inf_rows):
    ws = wb.create_sheet("Inf_Adicionais")
    headers = ["Ano", "Mes", "TotalFunc", "PCD_Ativos", "PCD_%", "Aprendiz_Ativos"]
    _table(ws, headers, inf_rows, [6, 6, 12, 12, 10, 16])


def build_leiame(wb):
    ws = wb.create_sheet("Leia-me")
    ws.column_dimensions["A"].width = 115
    ws.sheet_view.showGridLines = False
    linhas = [
        ("Como usar este modelo de dados", F_H1),
        ("", F_CELL),
        ("1. Aba Config: defina titulo, area e o PERIODO (Ano, Mes, Tipo).", F_CELL),
        ("   - Tipo 'Mensal' = fecha o mes escolhido.", F_CELL),
        ("   - Tipo 'Semestral' = 1o sem (mes<=6) ou 2o sem (mes>=7) do ano.", F_CELL),
        ("   - Tipo 'Anual' = jan a dez do ano.", F_CELL),
        ("", F_CELL),
        ("2. Aba Base_Colaboradores: cole 1 LINHA POR COLABORADOR.", F_CELL),
        ("   - Inclua ativos + admitidos/desligados dos ultimos 12 meses.", F_CELL),
        ("   - DataDesligamento vazia = colaborador ATIVO.", F_CELL),
        ("   - Daqui a automacao calcula HC, genero, piramide, tempo de casa,", F_CELL),
        ("     geracao, raca, span, PCD, aprendiz, admitidos, deslig. e turnover.", F_CELL),
        ("", F_CELL),
        ("3. Aba Custo_Rescisorio: 1 linha por desligado (Diretoria / Area / Custo).", F_CELL),
        ("", F_CELL),
        ("4. Aba Absenteismo: 1 linha por mes. Absenteismo % = Horas Perdidas /", F_CELL),
        ("   Horas Previstas. Aceita horas (numero) ou duracao.", F_CELL),
        ("", F_CELL),
        ("5. Aba Terceiros: 1 linha por terceiro (entrada/saida/parceiro/orcamento).", F_CELL),
        ("   O Grau de Terceirizacao (GTER) = terceiros / (proprios + terceiros).", F_CELL),
        ("", F_CELL),
        ("6. Aba Historico_Mensal: OPCIONAL. Por padrao a evolucao e calculada da", F_CELL),
        ("   base. HC_Projetado alimenta a linha 'Projetado' do grafico de HC.", F_CELL),
        ("", F_CELL),
        ("7. Aba Inf_Adicionais: OPCIONAL. Referencia da cota legal PCD/Aprendiz.", F_CELL),
        ("", F_CELL),
        ("8. Apague as linhas de EXEMPLO e rode:  py -3 gerar_apresentacao.py", F_CELL),
        ("", F_CELL),
        ("Valores aceitos (use os menus suspensos):", F_HINT),
        ("Genero: M, F  |  Motivo: Voluntario, Involuntario  |  Lider/PCD: S, N", F_HINT),
        ("Situacao: Ativo, Afastado, Inativo  |  Vinculo: CLT, Estagio, Aprendiz", F_HINT),
        ("Raca: Branca, Parda, Preta, Amarela, Indigena, Nao Informado", F_HINT),
        ("Terceiros -> Orcamento: CAPEX, OPEX.  Datas no formato DD/MM/AAAA.", F_HINT),
    ]
    for i, (txt, font) in enumerate(linhas, start=1):
        c = ws.cell(i, 1, txt)
        c.font, c.alignment = font, LEFT


def main():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    base = _exemplo_base()
    ativos = [r for r in base if r[5] is None]
    lideres = sum(1 for r in ativos if r[11] == "S")
    span_base = round((len(ativos) - lideres) / lideres, 1) if lideres else 1.5
    hist = _exemplo_historico(len(ativos), span_base)
    resc = _exemplo_rescisorio(base)
    absent = _exemplo_absenteismo(len(ativos))
    terceiros = _exemplo_terceiros()
    inf = _exemplo_inf(base)

    build_leiame(wb)
    build_config(wb)
    build_base(wb, base)
    build_historico(wb, hist)
    build_rescisorio(wb, resc)
    build_absenteismo(wb, absent)
    build_terceiros(wb, terceiros)
    build_inf_adicionais(wb, inf)

    wb.save(OUT)
    print(f"OK -> {OUT} ({len(wb.sheetnames)} abas, {len(base)} colaboradores de exemplo)")


if __name__ == "__main__":
    main()
