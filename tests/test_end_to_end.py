"""Smoke test do fluxo completo: planilha-modelo -> calculo -> PowerPoint."""
import os
import sys

from pptx import Presentation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

import gerar_apresentacao  # noqa: E402
import gerar_modelo_excel  # noqa: E402
from hr_deck import dados  # noqa: E402


def test_planilha_vira_apresentacao(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    gerar_modelo_excel.main()
    assert os.path.exists(gerar_modelo_excel.OUT)

    meta, slides = dados.montar_slides(gerar_modelo_excel.OUT)
    assert meta["periodo"]
    assert len(slides) > 10

    out = tmp_path / "saida.pptx"
    gerar_apresentacao.build_presentation(meta, slides, str(out))

    prs = Presentation(str(out))
    assert len(prs.slides) == len(slides) + 2  # capa + conteudo + fechamento
    charts = sum(1 for s in prs.slides for sh in s.shapes if sh.has_chart)
    assert charts > 0, "os graficos devem ser nativos (editaveis), nao imagens"


def test_sem_planilha_usa_dados_de_exemplo(tmp_path):
    out = tmp_path / "exemplo.pptx"
    slides = gerar_apresentacao.slides_from_spec()
    gerar_apresentacao.build_presentation(dict(gerar_apresentacao.spec.META), slides, str(out))
    assert out.exists()
